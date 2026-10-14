#!/usr/bin/env python3
"""
Plantilla SINTÉTICA para probar build-slides.py sin la plantilla real.
No es el branding de AIforDevelopers: usa el tema por defecto de python-pptx y formas genéricas
con nombres, para ejercitar clonado, tarjetas, imágenes, pie, verificación y render.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

out = sys.argv[1]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
ACC = RGBColor(0x4F, 0x81, 0xBD)  # accent1 del tema por defecto
DK = RGBColor(0x00, 0x00, 0x00)
LT = RGBColor(0xFF, 0xFF, 0xFF)

def box(s, name, x, y, w, h, text, size, color=DK, fill=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)) if fill else s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    tf = sh.text_frame; tf.word_wrap = True
    tf.text = text
    for r in tf.paragraphs[0].runs:
        r.font.size = Pt(size); r.font.color.rgb = color
    return sh

def footer(s, n):
    box(s, 'Pie', 0.4, 7.0, 4, 0.35, 'AI for Developers | 30X', 10)
    box(s, 'Número de diapositiva', 12.4, 7.0, 0.6, 0.35, str(n), 10)

s = prs.slides.add_slide(blank); box(s, 'Espera', 1, 3, 11, 1.2, 'Estamos por comenzar', 44)
s = prs.slides.add_slide(blank); box(s, 'Título', 1, 2.5, 11, 1.2, 'Estación 1', 48); box(s, 'Subtítulo', 1, 3.8, 11, 0.8, 'Introducción y demostración', 24)
s = prs.slides.add_slide(blank); footer(s, 3); box(s, 'Título', 0.5, 0.4, 12, 0.8, 'Agenda de hoy', 32)
for i in range(4):
    x = 0.5 + i * 3.1
    box(s, f'Tarjeta {i+1} fondo', x, 1.6, 2.9, 3.6, '', 12, fill=ACC)
    box(s, f'Tarjeta {i+1} número', x + 0.2, 1.8, 1, 0.6, f'0{i+1}', 28, LT)
    box(s, f'Tarjeta {i+1} título', x + 0.2, 2.5, 2.5, 0.9, 'Empieza por el producto', 18, LT)
    box(s, f'Tarjeta {i+1} texto', x + 0.2, 3.5, 2.5, 1.5, 'La ruta, el Product Vision Board y las reglas de la demo', 13, LT)
box(s, 'Cierre', 0.5, 6.2, 12, 0.6, 'Hoy la demo recorre toda esta ruta en 30 minutos.', 14)
s = prs.slides.add_slide(blank); footer(s, 5); box(s, 'Número grande', 1, 1.5, 3, 2, '01', 96, ACC); box(s, 'Título', 1, 3.6, 11, 1, 'Empieza por el producto', 40); box(s, 'Subtítulo', 1, 4.7, 11, 0.7, 'Product Vision Board · el entregable de hoy', 20)
s = prs.slides.add_slide(blank); footer(s, 11); box(s, 'Título', 0.5, 0.4, 12, 0.8, '¿Qué acaba de pasar?', 32)
for i in range(6):
    x = 0.5 + (i % 3) * 4.1; y = 1.5 + (i // 3) * 2.3
    box(s, f'Tarjeta {i+1} fondo', x, y, 3.9, 2.1, '', 12, fill=ACC)
    box(s, f'Tarjeta {i+1} número', x + 0.2, y + 0.15, 0.6, 0.5, str(i + 1), 24, LT)
    box(s, f'Tarjeta {i+1} título', x + 0.9, y + 0.15, 2.8, 0.5, 'Contexto antes que código', 16, LT)
    box(s, f'Tarjeta {i+1} texto', x + 0.2, y + 0.75, 3.5, 1.2, 'PVB → PRD → plan. El agente no escribió una línea hasta tener reglas.', 12, LT)
box(s, 'Cierre', 0.5, 6.3, 12, 0.5, 'Hoy fue la versión exprés.', 14)
s = prs.slides.add_slide(blank); footer(s, 7); box(s, 'Título', 0.5, 0.4, 12, 0.8, '¿Por qué importa más con un agente?', 32)
box(s, 'Columna izquierda', 0.5, 1.5, 5.9, 4.5, 'Precios, envío y descuentos inventados', 16); box(s, 'Columna derecha', 6.9, 1.5, 5.9, 4.5, 'Reglas explícitas: $99, gratis desde $600', 16)
box(s, 'Cierre', 0.5, 6.3, 12, 0.5, 'El agente amplifica lo que le das.', 14)
s = prs.slides.add_slide(blank); footer(s, 12); box(s, 'Título', 0.5, 0.4, 12, 0.8, 'Tu entregable', 32)
box(s, 'Caja imagen 1', 0.5, 1.4, 3.2, 4.8, 'imagen', 12, fill=ACC); box(s, 'Caja imagen 2', 3.9, 1.4, 3.2, 4.8, 'imagen', 12, fill=ACC)
box(s, 'Texto', 7.4, 1.4, 5.4, 4.8, 'Está bien hecho si la visión cabe en una frase', 16); box(s, 'Cierre', 0.5, 6.3, 12, 0.5, 'Entrega antes de la E2.', 14)
s = prs.slides.add_slide(blank); footer(s, 13); box(s, 'Signo', 1, 1.5, 11, 3, '?', 160, ACC); box(s, 'Subtítulo', 1, 4.8, 11, 0.8, 'Preguntas', 28)
prs.save(out)
print(out)
