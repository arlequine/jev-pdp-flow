#!/usr/bin/env python3
"""
npm run slides  [-- --template ruta.pptx] [--out docs/slides/charla-jev.pptx] [--no-render]

Genera el deck de la charla a partir de la plantilla de AIforDevelopers:

1. Abre la plantilla y usa SOLO su branding: master, layouts, tema, logos y las formas de sus
   diapositivas de ejemplo (que se clonan tal cual; no se dibujan formas a mano para imitarlas).
2. Lee el contenido de los bloques ```slide de docs/charla.md (única fuente del deck).
3. Reemplaza los marcadores {{…}} con cifras de metrics/ (npm run cost -- --json y before-after.json).
4. Borra todas las diapositivas originales de la plantilla.
5. Verifica: ningún texto de la plantilla queda en el deck, ningún texto desborda su caja,
   solo colores y fuentes del tema; y renderiza a PNG con LibreOffice para revisión visual.

Qué forma de qué diapositiva de ejemplo recibe qué campo está en docs/brand/patterns.json
(se escribe una vez, después de revisar el resumen de `npm run slides:brand`).
"""
import argparse, copy, hashlib, json, re, shutil, subprocess, sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
ap = argparse.ArgumentParser()
ap.add_argument('--template', default=str(ROOT / 'docs/brand/AIforDevelopers_E1_Introduccion_Demostracion.pptx'))
ap.add_argument('--patterns', default=str(ROOT / 'docs/brand/patterns.json'))
ap.add_argument('--brand', default=str(ROOT / 'docs/brand/brand.json'))
ap.add_argument('--charla', default=str(ROOT / 'docs/charla.md'))
ap.add_argument('--out', default=str(ROOT / 'docs/slides/charla-jev.pptx'))
ap.add_argument('--metrics', default=str(ROOT / 'metrics'))
ap.add_argument('--no-render', action='store_true')
args = ap.parse_args()

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
NS = {'a': A, 'r': R, 'p': P}

for f in (args.template, args.patterns, args.brand):
    if not Path(f).exists():
        sys.exit(f'Falta {f}. Orden: copiar la plantilla → npm run slides:brand → escribir patterns.json → npm run slides.')

brand = json.loads(Path(args.brand).read_text())
patterns = json.loads(Path(args.patterns).read_text())

# ───────────────────────── 1. Contenido desde docs/charla.md ─────────────────────────

def parse_slides(md: str):
    out = []
    for block in re.findall(r'```slide\n(.*?)```', md, re.S):
        spec, key = {}, None
        for line in block.splitlines():
            if not line.strip():
                continue
            m = re.match(r'^([a-záéíóúñ0-9-]+):\s*(.*)$', line)
            if m and not line.startswith('- '):
                key, val = m.group(1), m.group(2).strip()
                spec[key] = val if val else []
            elif line.startswith('- ') and key:
                if not isinstance(spec.get(key), list):
                    spec[key] = []
                spec[key].append(line[2:].strip())
        if 'tarjetas' in spec:
            spec['tarjetas'] = [dict(zip(('num', 'title', 'text'), [x.strip() for x in c.split('|')])) for c in spec['tarjetas']]
        out.append(spec)
    return out

slides_spec = parse_slides(Path(args.charla).read_text())
if not slides_spec:
    sys.exit('No encontré bloques ```slide en docs/charla.md')

# ───────────────────────── 2. Cifras desde metrics/ ─────────────────────────

def load_numbers():
    nums, missing = {}, []
    try:
        cost = json.loads(subprocess.run(['npx', 'tsx', 'packages/jev-flow/src/cost/cli.ts', '--json'], cwd=ROOT, capture_output=True, text=True, check=True).stdout.strip().splitlines()[-1])
    except Exception as e:  # noqa: BLE001
        print(f'aviso: no pude leer costos ({e})')
        cost = {}
    def usd(v):
        if v is None: return None
        return '$0' if v == 0 else (f'${v:.6f}'.rstrip('0') if v < 0.01 else f'${v:.4f}'.rstrip('0').rstrip('.'))
    for p in cost.get('prioritize', []):
        key = {'jev': 'jev', 'pedirle-a-sonnet': 'sonnet', 'pedirle-a-opus': 'opus'}.get(p['method'])
        if key:
            nums[f'prioritize.{key}.cost'] = usd(p.get('costUsd'))
            nums[f'prioritize.{key}.passed'] = None if p.get('passed') is None else str(p['passed'])
    ref = cost.get('refactor') or {}
    if ref.get('rows'):
        nums['refactor.cost'] = usd(ref.get('totalUsd'))
        nums['refactor.opus'] = usd(ref.get('opusUsd'))
    for r in cost.get('review', []):
        key = {'jev': 'jev', 'opus': 'opus', 'baseline': 'baseline'}.get(r['method'])
        if key:
            nums[f'review.{key}.cost'] = usd(r.get('costUsd'))
            nums[f'review.{key}.exact'] = f"{r['exact']}/{r['total']}"
    for k, v in (cost.get('prices') or {}).items():
        if k.startswith('jev'):
            nums['prices.jev.input'] = f"{v['input']:g}"
    ba_f = Path(args.metrics) / 'before-after.json'
    if ba_f.exists():
        ba = json.loads(ba_f.read_text())
        b, a = ba['before'], ba['after']
        nums['ba.requests'] = f"{b['apiRequests']} → {a['apiRequests']}"
        nums['ba.complete'] = f"{b['timeToCompleteMs']:,} ms → {a['timeToCompleteMs']:,} ms"
        nums['ba.price'] = f"{b['timeToPriceMs']:,} ms → {a['timeToPriceMs']:,} ms"
        nums['ba.a11y'] = f"{b['a11yScore']} → {a['a11yScore']}"
        nums['ba.lines'] = f"{b['componentLines']} → {a['componentLines']}"
        nums['ba.bundle'] = f"{b['bundleGzipKb']} KB → {a['bundleGzipKb']} KB"
    return nums

NUMBERS = load_numbers()
MISSING = set()
def fill(text: str) -> str:
    def rep(m):
        v = NUMBERS.get(m.group(1))
        if v is None:
            MISSING.add(m.group(1))
            return 's/d'
        return v
    return re.sub(r'\{\{\s*([\w.]+)\s*\}\}', rep, text)

def deep_fill(v):
    if isinstance(v, str): return fill(v)
    if isinstance(v, list): return [deep_fill(x) for x in v]
    if isinstance(v, dict): return {k: deep_fill(x) for k, x in v.items()}
    return v

slides_spec = [deep_fill(s) for s in slides_spec]

# ───────────────────────── 3. Construcción ─────────────────────────

prs = Presentation(args.template)
originals = list(prs.slides)
n_original = len(originals)

def set_text(shape, value, bullets=False):
    """Reemplaza el texto conservando el formato del primer párrafo/run de la plantilla."""
    tf = shape.text_frame
    lines = value if isinstance(value, list) else [value]
    paras = tf.paragraphs
    proto_p = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    def write(p_el, text):
        runs = p_el.findall('a:r', NS)
        for extra in runs[1:]:
            p_el.remove(extra)
        for br in p_el.findall('a:br', NS):
            p_el.remove(br)
        for fld in p_el.findall('a:fld', NS):
            p_el.remove(fld)
        if not runs:
            r = etree.SubElement(p_el, f'{{{A}}}r')
            etree.SubElement(r, f'{{{A}}}t')
            end = p_el.find('a:endParaRPr', NS)
            if end is not None:
                p_el.remove(r); end.addprevious(r)
            runs = [r]
        runs[0].find('a:t', NS).text = text
    write(first._p, lines[0] if lines else '')
    anchor = first._p
    for text in lines[1:]:
        p_new = copy.deepcopy(proto_p)
        write(p_new, text)
        anchor.addnext(p_new)
        anchor = p_new

def clone_slide(src_index: int):
    """Nueva diapositiva con el layout de la de ejemplo y copias exactas de sus formas (incluidas imágenes)."""
    src = originals[src_index - 1]
    dst = prs.slides.add_slide(src.slide_layout)
    for ph in list(dst.placeholders):
        ph._element.getparent().remove(ph._element)
    rid_map = {}
    for rel in src.part.rels.values():
        if rel.reltype in (RT.IMAGE, RT.MEDIA, RT.VIDEO) or rel.reltype.endswith('/hdphoto'):
            rid_map[rel.rId] = dst.part.relate_to(rel.target_part, rel.reltype)
        elif rel.reltype == RT.HYPERLINK:
            rid_map[rel.rId] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
    for el in src.shapes._spTree.iterchildren():
        tag = etree.QName(el).localname
        if tag in ('nvGrpSpPr', 'grpSpPr'):
            continue
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr in (f'{{{R}}}embed', f'{{{R}}}link', f'{{{R}}}id'):
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        dst.shapes._spTree.append(new)
    # Fondo propio de la diapositiva, si lo tiene.
    bg = src._element.find('p:cSld/p:bg', NS)
    if bg is not None:
        dst._element.find('p:cSld', NS).insert(0, copy.deepcopy(bg))
    return dst

def all_shapes(container):
    for sh in container:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from all_shapes(sh.shapes)

def by_name(slide, name):
    return next((s for s in all_shapes(slide.shapes) if s.name == name), None)

def remove(shape):
    shape._element.getparent().remove(shape._element)

def add_picture_in(slide, box, path: Path):
    """Inserta una imagen ajustada (sin deformar) dentro de la caja `box` de la plantilla."""
    from PIL import Image
    w, h = Image.open(path).size
    bw, bh = box.width, box.height
    scale = min(bw / w, bh / h)
    pw, ph = int(w * scale), int(h * scale)
    pic = slide.shapes.add_picture(str(path), box.left + (bw - pw) // 2, box.top + (bh - ph) // 2, pw, ph)
    pic.name = f'Imagen {path.stem}'
    return pic

def from_layout(cfg, spec):
    layout = prs.slide_layouts[cfg['layout']]
    s = prs.slides.add_slide(layout)
    for ph in s.placeholders:
        field = cfg.get('placeholders', {}).get(str(ph.placeholder_format.idx))
        if field and spec.get(field):
            set_text(ph, spec[field])
        elif ph.has_text_frame:
            remove(ph)
    return s

def build(spec, number):
    pat = spec.get('patrón')
    cfg = patterns.get(pat)
    if not cfg:
        sys.exit(f'patrón "{pat}" no está en patterns.json')
    s = from_layout(cfg, spec) if 'layout' in cfg else clone_slide(cfg['clone'])
    used = set()
    # Campos simples → formas por nombre.
    for shape_name, field in cfg.get('fill', {}).items():
        sh = by_name(s, shape_name)
        if sh is None:
            sys.exit(f'patterns.json: no existe la forma "{shape_name}" en el patrón {pat}')
        val = spec.get(field)
        if val in (None, [], ''):
            remove(sh)
        else:
            set_text(sh, val)
            used.add(shape_name)
    # Tarjetas: grupos (o conjuntos de formas) que se repiten.
    if 'cards' in cfg:
        cards = spec.get('tarjetas', [])
        slots = cfg['cards']['slots']
        if len(cards) > len(slots):
            sys.exit(f'El patrón {pat} de la plantilla tiene {len(slots)} tarjetas y la diapositiva pide {len(cards)}')
        for i, slot in enumerate(slots):
            for shape_name, field in slot.items():
                sh = by_name(s, shape_name)
                if sh is None:
                    continue
                if i < len(cards) and cards[i].get(field):
                    set_text(sh, cards[i][field])
                    used.add(shape_name)
                else:
                    remove(sh)
            for extra in cfg['cards'].get('decor', [[]] * len(slots))[i] if i >= len(cards) else []:
                sh = by_name(s, extra)
                if sh is not None:
                    remove(sh)
    # Imágenes dentro de cajas de la plantilla.
    for box_name, field in cfg.get('images', {}).items():
        box = by_name(s, box_name)
        img = spec.get(field)
        if box is None:
            continue
        if img and (ROOT / img).exists():
            add_picture_in(s, box, ROOT / img)
            if cfg.get('remove_image_boxes', True):
                remove(box)
        elif img:
            MISSING.add(f'imagen:{img}')
            set_text(box, f'Falta {img}') if box.has_text_frame else None
        else:
            remove(box)
    # Número de diapositiva del pie (texto fijo en la plantilla).
    if cfg.get('page_number') and by_name(s, cfg['page_number']) is not None:
        set_text(by_name(s, cfg['page_number']), str(number))
        used.add(cfg['page_number'])
    for name in cfg.get('drop', []):
        sh = by_name(s, name)
        if sh is not None:
            remove(sh)
    if spec.get('notas'):
        s.notes_slide.notes_text_frame.text = spec['notas'] if isinstance(spec['notas'], str) else '\n'.join(spec['notas'])
    return s

for i, spec in enumerate(slides_spec, start=1):
    build(spec, i)

# Borra las diapositivas originales de la plantilla (y sus partes).
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[:n_original]:
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

out = Path(args.out)
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(out)
print(f'{out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}: {len(slides_spec)} diapositivas')

# ───────────────────────── 4. Verificación ─────────────────────────

problems = []
deck = Presentation(str(out))
FORBIDDEN = [r'\bGrano\b', r'\bPVB\b', r'Product Vision Board', r'Andr[eé]s Bilbao', r'\bEstaci[oó]n\b', r'\bE(?:[1-9]|1[01])\b', r'\bCOT\b', r'Huila', r'Chiapas', r'\$837', r'BIENVENIDA10', r'\bRender\b', r'\bVercel\b', r'\bISB\b']
# "?" y "Preguntas" son parte del patrón de cierre que se pide conservar; no cuentan como contenido de la plantilla.
ALLOWED = {'?', 'Preguntas'}
template_hashes = set(brand.get('template_texts_sha256', [])) - {hashlib.sha256(t.encode()).hexdigest() for t in ALLOWED}
theme_colors = {c.upper().lstrip('#') for c in brand['theme']['colors'].values()}
theme_fonts = {brand['theme']['fonts']['major'], brand['theme']['fonts']['minor']}

def est_overflow(sh):
    """Estimación conservadora: líneas × alto de línea contra el alto de la caja."""
    if not sh.has_text_frame or not sh.text_frame.text.strip() or not sh.width or not sh.height:
        return None
    tf = sh.text_frame
    bodyPr = tf._txBody.find('a:bodyPr', NS)
    if bodyPr is not None and (bodyPr.find('a:normAutofit', NS) is not None or bodyPr.find('a:spAutoFit', NS) is not None):
        pass  # aun con autoajuste, medimos: el autoajuste no se aplica hasta abrir el archivo
    ins = lambda k, d: int(bodyPr.get(k, d)) if bodyPr is not None and bodyPr.get(k) else d
    width = sh.width - ins('lIns', 91440) - ins('rIns', 91440)
    height = sh.height - ins('tIns', 45720) - ins('bIns', 45720)
    total = 0
    for p in tf.paragraphs:
        size = next((r.font.size.pt for r in p.runs if r.font.size), None) or 18
        chars_per_line = max(1, int(width / Emu(Pt(size * 0.5))))
        lines = max(1, -(-len(p.text) // chars_per_line))
        total += lines * Pt(size * 1.2)
    return total > height * 1.05

for n, s in enumerate(deck.slides, start=1):
    for sh in all_shapes(s.shapes):
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t and hashlib.sha256(t.encode()).hexdigest() in template_hashes:
                problems.append(f'D{n}: quedó texto de la plantilla en "{sh.name}"')
            for pat in FORBIDDEN:
                if re.search(pat, t):
                    problems.append(f'D{n}: texto prohibido ({pat}) en "{sh.name}": {t[:50]!r}')
            if est_overflow(sh):
                problems.append(f'D{n}: posible desborde en "{sh.name}" ({len(t)} caracteres)')
            for el in sh._element.iter(f'{{{A}}}srgbClr'):
                if el.get('val', '').upper() not in theme_colors:
                    problems.append(f'D{n}: color fuera del tema #{el.get("val")} en "{sh.name}"')
            for el in sh._element.iter(f'{{{A}}}latin'):
                face = el.get('typeface', '')
                if face and not face.startswith('+') and face not in theme_fonts:
                    problems.append(f'D{n}: fuente fuera del tema "{face}" en "{sh.name}"')

if MISSING:
    print('Sin dato (aparecen como "s/d"): ' + ', '.join(sorted(MISSING)))
if problems:
    print('\nProblemas:')
    for p in sorted(set(problems)):
        print('  - ' + p)
else:
    print('Verificación: sin texto de la plantilla, sin desbordes estimados, solo colores y fuentes del tema.')

# ───────────────────────── 5. Render a PNG ─────────────────────────

if not args.no_render:
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        print('No encontré LibreOffice: omito el render a PNG.')
    else:
        png_dir = out.parent / 'png'
        if png_dir.exists():
            shutil.rmtree(png_dir)
        png_dir.mkdir(parents=True)
        subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(png_dir), str(out)], check=True, capture_output=True)
        pdf = png_dir / (out.stem + '.pdf')
        if shutil.which('pdftoppm'):
            subprocess.run(['pdftoppm', '-png', '-r', '60', str(pdf), str(png_dir / 'D')], check=True)
            print(f'PNG de revisión en {png_dir.relative_to(ROOT) if png_dir.is_relative_to(ROOT) else png_dir}/ (uno por diapositiva)')
        else:
            print(f'PDF de revisión en {pdf}')

sys.exit(1 if problems else 0)
